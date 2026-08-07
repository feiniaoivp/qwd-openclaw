#!/usr/bin/env python3
"""
TTM 分析报告生成器 — 补齐季报做TTM指标（滚动ROE、滚动F-Score）
输出: HTML交互报告 + Excel报告
"""
import sys; sys.path.insert(0, '.')
from ttm_core import analyze_single_deep, analyze_stocks
import pandas as pd
import numpy as np
from datetime import datetime
import json
import warnings; warnings.filterwarnings('ignore')

STOCK_POOL = [
    ('600160','巨化股份'),('300827','上能电气'),('600346','恒力石化'),
    ('000708','中信特钢'),('300748','金力永磁'),('002413','雷科防务'),
    ('601061','中信金属'),('000157','中联重科'),('603308','应流股份'),
    ('601865','福莱特'),('600660','福耀玻璃'),('002318','久立特材'),
    ('002335','科华数据'),('605566','福莱蒽特'),('601995','中金公司'),
    ('600030','中信证券'),('601066','中信建投'),('000987','越秀资本'),
    ('601100','恒立液压'),('600570','恒生电子'),('300124','汇川技术'),
    ('300719','安达维尔'),('002466','天齐锂业'),('300014','亿纬锂能'),
    ('600036','招商银行'),('300285','国瓷材料'),('688981','中芯国际'),
    ('600584','长电科技'),('002156','通富微电'),('603601','再升科技'),
    ('605123','派克新材'),('002149','西部材料'),('003009','中天火箭'),
    ('603678','火炬电子'),('300337','银邦股份'),('603596','伯特利'),
    ('688433','华曙高科'),('300762','上海瀚讯'),('600343','航天动力'),
]

def fmt(v, decimals=2, prefix='', suffix=''):
    if v is None or (isinstance(v, float) and np.isnan(v)): return 'N/A'
    return f'{prefix}{v:,.{decimals}f}{suffix}'

def _fdot(fd, i):
    p = fd.split('|') if fd else []
    l = p[i] if i < len(p) else f'F{i+1}?'
    c = 'on' if i < len(p) and '+' in p[i] else 'off'
    return f'<span class="f-dot {c}" title="F{i+1}">{l}</span>'

def _card(r):
    c = r.get('股票代码', '')
    n = r.get('股票名称', c)
    roe, fs = r.get('TTM ROE(%)'), r.get('F-Score',0)
    npv, rv, cfv = r.get('TTM归母净利润(亿)'), r.get('TTM营业总收入(亿)'), r.get('TTM经营现金流(亿)')
    eq = r.get('净资产(亿)'); liab = r.get('资产负债率(%)')
    crv = r.get('流动比率'); gmv = r.get('毛利率(%)'); atv = r.get('总资产周转率(TTM)')
    rd = r.get('报告期',''); fd = r.get('F_明细','')
    rc = '#22c55e' if roe is not None and roe>=15 else ('#eab308' if roe is not None and roe>=8 else '#ef4444')
    fc = '#22c55e' if fs>=7 else ('#eab308' if fs>=4 else '#ef4444')
    dots = ''.join(_fdot(fd,i) for i in range(9))
    return f'''<div class="card" onclick="toggleDetail('{c}')">
<div class="card-header"><span class="stock-title">{n} <span class="stock-code">{c}</span></span><span class="report-date">📅 {rd}</span></div>
<div class="card-body">
<div class="metrics-row">
<div class="metric"><div class="metric-label">TTM ROE</div><div class="metric-value" style="color:{rc}">{fmt(roe,1,suffix='%')}</div></div>
<div class="metric"><div class="metric-label">F-Score</div><div class="metric-value" style="color:{fc}">{fs}<span class="sub">/9</span></div></div>
<div class="metric"><div class="metric-label">TTM净利润</div><div class="metric-value">{fmt(npv,1,suffix='亿')}</div></div>
<div class="metric"><div class="metric-label">TTM营收</div><div class="metric-value">{fmt(rv,1,suffix='亿')}</div></div>
</div>
<div class="f-score-bar">{dots}</div>
<div id="detail-{c}" class="detail-row" style="display:none"><table class="detail-table">
<tr><td>资产负债率</td><td>{fmt(liab,1,suffix='%')}</td><td>流动比率</td><td>{fmt(crv,2)}</td></tr>
<tr><td>毛利率</td><td>{fmt(gmv,1,suffix='%')}</td><td>资产周转率</td><td>{fmt(atv,4)}</td></tr>
<tr><td>经营现金流</td><td>{fmt(cfv,1,suffix='亿')}</td><td>净资产</td><td>{fmt(eq,1,suffix='亿')}</td></tr>
</table></div></div></div>'''

def gen_html(result, all_ttms, path='ttm_report.html'):
    cards = ''.join(_card(r) for _,r in result.iterrows())
    td = {}
    for code, t in all_ttms.items():
        t = t.sort_values('报告期').tail(20)
        td[code] = {
            'dates': t['报告期'].tolist(),
            'roe': [None if pd.isna(v) else round(v,2) for v in t['TTM ROE(%)']],
            'fscore': [0 if pd.isna(v) else int(v) for v in t['F-Score']],
            'revenue': [None if pd.isna(v) else round(v,1) for v in t['TTM营业总收入(亿)']],
        }
    tj = json.dumps(td, ensure_ascii=False)
    
    hr = sum(1 for _,r in result.iterrows() if r.get('TTM ROE(%)') is not None and r['TTM ROE(%)'] >= 15)
    hf = sum(1 for _,r in result.iterrows() if r.get('F-Score',0) >= 7)
    total = len(result)
    
    cl, nl = result['股票代码'].tolist(), result['股票名称'].tolist()
    cb = ''.join(f'<button class="chart-btn {"active" if i==0 else ""}" onclick="switchStock(\'{c}\',this)">{n}</button>' for i,(c,n) in enumerate(list(zip(cl,nl))[:12]))
    dc = cl[0] if cl else ''
    
    html = f'''<!DOCTYPE html><html lang="zh-CN">
<head><meta charset="UTF-8"><meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>A股 TTM指标分析报告</title>
<script src="https://cdn.jsdelivr.net/npm/chart.js@4.4.0/dist/chart.umd.min.js"></script>
<style>
*{{margin:0;padding:0;box-sizing:border-box}}
body{{font-family:-apple-system,'PingFang SC','Microsoft YaHei',sans-serif;background:#0f172a;color:#e2e8f0;min-height:100vh}}
.header{{background:linear-gradient(135deg,#1e293b,#0f172a);padding:30px 20px;border-bottom:1px solid #334155}}
.header h1{{font-size:28px;font-weight:700;background:linear-gradient(135deg,#60a5fa,#a78bfa);-webkit-background-clip:text;-webkit-text-fill-color:transparent}}
.header .subtitle{{color:#94a3b8;margin-top:8px;font-size:14px}}
.stats-bar{{display:flex;gap:20px;padding:20px;background:#1e293b;border-bottom:1px solid #334155;flex-wrap:wrap}}
.stat-box{{flex:1;min-width:120px;text-align:center;padding:12px;background:#0f172a;border-radius:12px;border:1px solid #334155}}
.stat-box .num{{font-size:24px;font-weight:700}}
.stat-box .label{{font-size:12px;color:#94a3b8;margin-top:4px}}
.container{{max-width:1400px;margin:0 auto;padding:20px}}
.grid{{display:grid;grid-template-columns:repeat(auto-fill,minmax(380px,1fr));gap:16px}}
.card{{background:#1e293b;border-radius:16px;border:1px solid #334155;overflow:hidden;cursor:pointer;transition:all 0.2s}}
.card:hover{{border-color:#60a5fa;transform:translateY(-2px);box-shadow:0 8px 24px rgba(96,165,250,0.15)}}
.card-header{{display:flex;justify-content:space-between;align-items:center;padding:16px 20px;background:#0f172a;border-bottom:1px solid #334155}}
.stock-title{{font-size:16px;font-weight:600}}
.stock-code{{color:#64748b;font-size:13px;font-weight:400;margin-left:8px}}
.report-date{{color:#64748b;font-size:12px}}
.card-body{{padding:16px 20px}}
.metrics-row{{display:flex;gap:12px}}
.metric{{flex:1;text-align:center;padding:8px;background:#0f172a;border-radius:10px}}
.metric-label{{font-size:11px;color:#94a3b8;margin-bottom:4px}}
.metric-value{{font-size:20px;font-weight:700}}
.sub{{font-size:12px;color:#64748b;font-weight:400}}
.f-score-bar{{display:flex;gap:4px;margin-top:12px;justify-content:center}}
.f-dot{{display:inline-flex;align-items:center;justify-content:center;width:32px;height:22px;border-radius:4px;font-size:10px;font-weight:600}}
.f-dot.on{{background:#166534;color:#4ade80}}
.f-dot.off{{background:#7f1d1d;color:#f87171}}
.detail-row{{margin-top:12px;padding-top:12px;border-top:1px solid #334155}}
.detail-table{{width:100%;font-size:13px}}
.detail-table td{{padding:4px 8px}}
.detail-table td:nth-child(odd){{color:#94a3b8;text-align:right;width:35%}}
.detail-table td:nth-child(even){{color:#e2e8f0;text-align:left;width:15%}}
.chart-section{{margin-top:30px;background:#1e293b;border-radius:16px;border:1px solid #334155;padding:20px}}
.chart-controls{{display:flex;gap:8px;margin-bottom:16px;flex-wrap:wrap}}
.chart-btn{{padding:6px 16px;border-radius:8px;border:1px solid #334155;background:#0f172a;color:#94a3b8;cursor:pointer;font-size:13px;transition:all 0.2s}}
.chart-btn.active{{background:#1d4ed8;color:#fff;border-color:#1d4ed8}}
.chart-btn:hover{{border-color:#60a5fa}}
.chart-container{{position:relative;height:400px;width:100%}}
.search-box{{margin-bottom:16px}}
.search-box input{{width:100%;padding:10px 16px;border-radius:10px;border:1px solid #334155;background:#0f172a;color:#e2e8f0;font-size:14px;outline:none}}
.search-box input:focus{{border-color:#60a5fa}}
.sort-controls{{display:flex;gap:8px;margin-bottom:12px}}
.sort-btn{{padding:6px 14px;border-radius:6px;border:1px solid #334155;background:#0f172a;color:#94a3b8;cursor:pointer;font-size:12px;transition:all 0.2s}}
.sort-btn.active{{background:#334155;color:#e2e8f0}}
.footer{{text-align:center;padding:20px;color:#475569;font-size:12px}}
@media(max-width:640px){{.grid{{grid-template-columns:1fr}}.header h1{{font-size:22px}}}}
</style></head><body>
<div class="header"><h1>📊 A股 TTM指标分析报告</h1><div class="subtitle">滚动ROE · Piotroski F-Score · 补齐季报生成TTM指标</div><div class="subtitle" style="margin-top:4px">生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M')} ｜ 数据: 同花顺/东方财富</div></div>
<div class="stats-bar"><div class="stat-box"><div class="num" style="color:#60a5fa">{total}</div><div class="label">分析股票</div></div><div class="stat-box"><div class="num" style="color:#22c55e">{hr}</div><div class="label">高ROE(≥15%)</div></div><div class="stat-box"><div class="num" style="color:#a78bfa">{hf}</div><div class="label">高F-Score(≥7)</div></div></div>
<div class="container">
<div class="search-box"><input type="text" id="si" placeholder="🔍 搜索..." oninput="fc()"></div>
<div class="sort-controls"><button class="sort-btn active" onclick="sc('d',this)">📋 默认</button><button class="sort-btn" onclick="sc('r',this)">📈 ROE排序</button><button class="sort-btn" onclick="sc('f',this)">🏆 F-Score排序</button></div>
<div class="grid" id="sg">{cards}</div>
<div class="chart-section"><h3 style="margin-bottom:12px">📈 TTM ROE / F-Score 趋势</h3><div class="chart-controls" id="cc">{cb}</div><div class="chart-container"><canvas id="tc"></canvas></div></div>
</div>
<div class="footer">⚠️ 仅供研究参考，不构成投资建议 | 数据来源: 同花顺/东方财富(AkShare)</div>
<script>const TD={tj};let cc=null,cs='{dc}';
function fc(){{const q=document.getElementById('si').value.toLowerCase();document.querySelectorAll('.card').forEach(c=>c.style.display=c.textContent.toLowerCase().includes(q)?'':'none')}}
function td(id){{const e=document.getElementById('detail-'+id);if(e)e.style.display=e.style.display==='none'?'block':'none'}}
function sc(m,b){{
document.querySelectorAll('.sort-btn').forEach(x=>x.classList.remove('active'));if(b)b.classList.add('active');
const g=document.getElementById('sg'),cards=Array.from(g.children);
cards.sort((a,b)=>{{
if(m==='r'){{const ra=parseFloat(a.querySelectorAll('.metric-value')[0].textContent)||0,rb=parseFloat(b.querySelectorAll('.metric-value')[0].textContent)||0;return rb-ra}}
if(m==='f'){{const fa=parseInt(a.querySelectorAll('.metric-value')[1].textContent.split('/')[0])||0,fb=parseInt(b.querySelectorAll('.metric-value')[1].textContent.split('/')[0])||0;return fb-fa}}
return 0
}});cards.forEach(c=>g.appendChild(c))
}}
function sw(id,b){{cs=id;document.querySelectorAll('.chart-btn').forEach(x=>x.classList.remove('active'));if(b)b.classList.add('active');rc(id)}}
function rc(id){{
const d=TD[id];if(!d||!d.dates)return;if(cc)cc.destroy();
cc=new Chart(document.getElementById('tc').getContext('2d'),{{
type:'line',data:{{labels:d.dates.map(x=>x.slice(5)),datasets:[
{{label:'TTM ROE(%)',data:d.roe,borderColor:'#60a5fa',backgroundColor:'rgba(96,165,250,0.1)',fill:true,tension:0.3,pointRadius:3,pointHoverRadius:6,yAxisID:'y'}},
{{label:'F-Score',data:d.fscore,borderColor:'#a78bfa',backgroundColor:'rgba(167,139,250,0.1)',fill:true,tension:0.3,pointRadius:3,pointHoverRadius:6,yAxisID:'y1'}},
{{label:'TTM营收(亿)',data:d.revenue,borderColor:'#f59e0b',backgroundColor:'rgba(245,158,11,0.05)',fill:false,tension:0.3,pointRadius:2,borderDash:[5,5],yAxisID:'y2'}}
]}},options:{{
responsive:true,maintainAspectRatio:false,interaction:{{mode:'index',intersect:false}},
plugins:{{legend:{{labels:{{color:'#94a3b8',usePointStyle:true}}}}}},
scales:{{
x:{{ticks:{{color:'#64748b',maxRotation:45}},grid:{{color:'#1e293b'}}}},
y:{{type:'linear',display:true,position:'left',ticks:{{color:'#60a5fa'}},grid:{{color:'#1e293b'}},title:{{display:true,text:'ROE(%)',color:'#94a3b8'}}}},
y1:{{type:'linear',display:true,position:'right',ticks:{{color:'#a78bfa',stepSize:1}},grid:{{display:false}},min:0,max:9,title:{{display:true,text:'F-Score',color:'#94a3b8'}}}},
y2:{{type:'linear',display:false,position:'right'}}
}}
}}}}}}}});if(cs)rc(cs)
</script></body></html>'''
    
    with open(path, 'w', encoding='utf-8') as f:
        f.write(html)
    print(f'✅ HTML: {path} ({len(html):,} bytes)')

def export_excel(result, all_ttms, path='ttm_report.xlsx'):
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    cols = ['股票代码','股票名称','报告期','TTM归母净利润(亿)','TTM营业总收入(亿)',
            'TTM扣非净利润(亿)','TTM经营现金流(亿)','TTM ROE(%)','资产负债率(%)',
            '流动比率','毛利率(%)','总资产周转率(TTM)','净资产(亿)','F-Score','F_明细']
    s1 = result[[c for c in cols if c in result.columns]].copy()
    
    with pd.ExcelWriter(path, engine='openpyxl') as w:
        s1.to_excel(w, sheet_name='最新TTM汇总', index=False)
        ws = w.sheets['最新TTM汇总']
        hf = PatternFill(start_color='1e3a5f', end_color='1e3a5f', fill_type='solid')
        hft = Font(color='FFF', bold=True, size=11)
        thin = Border(left=Side('thin','334155'), right=Side('thin','334155'),
                      top=Side('thin','334155'), bottom=Side('thin','334155'))
        for c in ws[1]: c.fill=hf; c.font=hft; c.alignment=Alignment(horizontal='center'); c.border=thin
        cw = {'A':12,'B':14,'C':12,'D':18,'E':18,'F':18,'G':18,'H':14,'I':14,'J':12,'K':12,'L':18,'M':12,'N':10,'O':30}
        for col,wc in cw.items(): ws.column_dimensions[col].width = wc
        # ROE coloring (col H = 8)
        for r in range(2, ws.max_row + 1):
            c = ws.cell(row=r, column=8)
            if c.value is not None:
                try:
                    v = float(c.value)
                    if v >= 15:
                        c.fill = PatternFill(start_color='14532d', end_color='14532d', fill_type='solid')
                        c.font = Font(color='4ade80', bold=True)
                    elif v >= 8:
                        c.fill = PatternFill(start_color='422006', end_color='422006', fill_type='solid')
                        c.font = Font(color='facc15', bold=True)
                    else:
                        c.fill = PatternFill(start_color='450a0a', end_color='450a0a', fill_type='solid')
                        c.font = Font(color='f87171', bold=True)
                except:
                    pass
        # F-Score coloring (col N = 14)
        for r in range(2, ws.max_row + 1):
            c = ws.cell(row=r, column=14)
            if c.value is not None:
                try:
                    v = int(c.value)
                    if v >= 7:
                        c.fill = PatternFill(start_color='14532d', end_color='14532d', fill_type='solid')
                        c.font = Font(color='4ade80', bold=True)
                    elif v >= 4:
                        c.fill = PatternFill(start_color='422006', end_color='422006', fill_type='solid')
                        c.font = Font(color='facc15', bold=True)
                    else:
                        c.fill = PatternFill(start_color='450a0a', end_color='450a0a', fill_type='solid')
                        c.font = Font(color='f87171', bold=True)
                except:
                    pass
        
        for code, t in all_ttms.items():
            nm = result[result['股票代码']==code]['股票名称'].values
            nm = nm[0] if len(nm)>0 else code
            sn = f'{nm}_{code}'[:31]
            t2 = t.sort_values('报告期', ascending=False)
            ecols = ['报告期','季度类型','TTM归母净利润(亿)','TTM营业总收入(亿)','TTM经营现金流(亿)',
                     'TTM ROE(%)','资产负债率(%)','流动比率','毛利率(%)','总资产周转率(TTM)','F-Score','F_明细']
            ok = [c for c in ecols if c in t2.columns]
            t2[ok].to_excel(w, sheet_name=sn, index=False)
    print(f'✅ Excel: {path}')

def run(full=False):
    pool = STOCK_POOL if full else STOCK_POOL[:6]
    codes = [c for c,_ in pool]
    names = {c:n for c,n in pool}
    
    print(f'🚀 批量TTM分析 ({len(codes)}只)')
    print(f'   {" ".join(names.values())}')
    
    result = analyze_stocks(codes, names)
    if result.empty: print('❌ 无数据'); return
    
    print(f'\n📊 批量分析完成: {len(result)} 只')
    print(result[['股票名称','股票代码','报告期','TTM ROE(%)','F-Score']].to_string(index=False))
    
    # 排名
    print(f'\n🏆 ROE排名:')
    for _, r in result.sort_values('TTM ROE(%)', ascending=False, na_position='last').iterrows():
        fsv = r.get('F-Score',0)
        roe_v = r.get('TTM ROE(%)',0) or 0
        print(f'  {"⭐" if fsv>=7 else "✅" if fsv>=4 else "⚠️"} {r["股票名称"]}({r["股票代码"]}): ROE={roe_v:.1f}%  F={fsv}/9')
    
    print(f'\n🏆 F-Score排名:')
    for _, r in result.sort_values('F-Score', ascending=False, na_position='last').iterrows():
        roe_v = r.get('TTM ROE(%)',0) or 0
        fsv = r.get('F-Score',0)
        print(f'  {r["股票名称"]}({r["股票代码"]}): F={fsv}/9  ROE={roe_v:.1f}%')
    
    print(f'\n🔄 获取个股时序...')
    all_ttms = {}
    for i, code in enumerate(codes):
        nm = names.get(code, code)
        try:
            ttm,_ = analyze_single_deep(code)
            if not ttm.empty:
                all_ttms[code] = ttm
            print(f'  [{i+1}/{len(codes)}] {nm} {"✓" if code in all_ttms else "❌"}')
        except Exception as e:
            print(f'  [{i+1}/{len(codes)}] {nm} ❌ {e}')
    
    gen_html(result, all_ttms, 'ttm_report.html')
    export_excel(result, all_ttms, 'ttm_report.xlsx')
    print('\n✅ 全部完成!')

if __name__ == '__main__':
    run(full='--full' in sys.argv)
